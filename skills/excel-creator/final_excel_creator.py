#!/usr/bin/env python3
"""
Excel Creator for Budget Tracking

This script creates a comprehensive home budget tracker Excel spreadsheet with multiple sheets,
formulas, data validation, and formatting as specified in the budget tracking requirements.
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def create_budget_tracker():
    """Create a comprehensive budget tracker Excel file with all required sheets"""
    
    # Create workbook and worksheets
    wb = Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    # Create the 5 required sheets
    expense_sheet = wb.create_sheet(title="Expense Tracking")
    categories_sheet = wb.create_sheet(title="Expense Categories")
    savings_sheet = wb.create_sheet(title="Savings Goals")
    summary_sheet = wb.create_sheet(title="Monthly Summary")
    budget_sheet = wb.create_sheet(title="Budget Tracking")
    
    # Define styles
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    